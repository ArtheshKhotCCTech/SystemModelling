# Purpose: pins each phase 3 file reader on small synthetic fixtures (built in tmp_path, plus the
# hand-made PDFs in tests/fixtures/ingest) — located chunks of the right kind, tables that keep
# their headers, email threads split into dated messages, code with comments kept, CSV summarised
# not inlined, images never silently skipped, and a broken file reported unread with a reason.
import json
from pathlib import Path

import docx
import openpyxl
import pytest

from specalive.config import load_settings
from specalive.ingest.readers import ReadContext, detect_format, read_source

FIXTURES = Path(__file__).parent / "fixtures" / "ingest"


@pytest.fixture
def ctx(tmp_path):
    return ReadContext(settings=load_settings({"SPECALIVE_CACHE_DIR": str(tmp_path / "cache")}),
                       llm=None)


def texts(result):
    return [c.text for c in result.chunks]


# --- registry -----------------------------------------------------------------------------

@pytest.mark.parametrize("name, fmt", [
    ("a.pdf", "pdf"), ("a.docx", "docx"), ("a.xlsx", "xlsx"), ("a.eml", "eml"),
    ("a.txt", "text"), ("a.md", "markdown"), ("a.mo", "modelica"), ("a.puml", "plantuml"),
    ("a.json", "json"), ("a.csv", "csv"), ("a.png", "image"), ("a.JPG", "image"),
])
def test_format_by_extension(tmp_path, name, fmt):
    p = tmp_path / name
    p.write_bytes(b"x")
    assert detect_format(p) == fmt


@pytest.mark.parametrize("data, fmt", [
    (b"%PDF-1.4\n...", "pdf"),
    (b"\x89PNG\r\n\x1a\n....", "image"),
    (b"@startuml\nA --> B\n@enduml\n", "plantuml"),
    (b"From: a@b.example\nSubject: hi\n\nbody\n", "eml"),
    (b"plain words\n", "text"),
    (b"\x00\x01\x02\xff\xfe binary", "unknown"),
])
def test_format_by_content_sniffing_without_extension(tmp_path, data, fmt):
    p = tmp_path / "noext"
    p.write_bytes(data)
    assert detect_format(p) == fmt


def test_unknown_format_is_unread_with_reason(tmp_path, ctx):
    p = tmp_path / "blob.bin"
    p.write_bytes(b"\x00\x01\x02\xff")
    r = read_source(p, "src_blob", ctx)
    assert r.status == "unread"
    assert "unsupported" in r.reason
    assert r.chunks == []


def test_reader_exception_becomes_unread_not_a_crash(tmp_path, ctx):
    p = tmp_path / "bad.json"
    p.write_text("{not json", encoding="utf-8")
    r = read_source(p, "src_bad", ctx)
    assert r.status == "unread"
    assert r.reason


# --- pdf ----------------------------------------------------------------------------------

def test_pdf_chunks_are_located_by_page_and_line(ctx):
    r = read_source(FIXTURES / "two_pages.pdf", "src_pdf", ctx)
    assert r.format == "pdf" and r.status == "read"
    by_text = {c.text: c for c in r.chunks}
    item = next(c for c in r.chunks if c.text.startswith("WR-F-001"))
    assert item.locator.startswith("page 1, line")
    assert "WR-F-002" not in item.text  # an id-led item is its own chunk
    page2 = [c for c in r.chunks if c.locator.startswith("page 2")]
    assert any("controller reads the level sensor" in c.text for c in page2)
    assert all(c.kind == "prose" and c.source_id == "src_pdf" for c in r.chunks)
    assert "The widget shall hold water." in "\n".join(by_text)


def test_pdf_headings_start_new_chunks(ctx):
    r = read_source(FIXTURES / "two_pages.pdf", "src_pdf", ctx)
    starts = [c.text.splitlines()[0] for c in r.chunks]
    assert "1. Scope" in starts and "2. Setpoints" in starts and "3. Interfaces" in starts


def test_corrupt_pdf_is_unread_with_reason(ctx):
    r = read_source(FIXTURES / "corrupt.pdf", "src_bad", ctx)
    assert r.status == "unread"
    assert r.reason and r.chunks == []


# --- docx ---------------------------------------------------------------------------------

@pytest.fixture
def docx_file(tmp_path):
    d = docx.Document()
    d.add_paragraph("PUMP CONTROL DESIGN NOTE")
    meta = d.add_table(rows=2, cols=2)
    meta.cell(0, 0).text, meta.cell(0, 1).text = "Document", "PCN-07"
    meta.cell(1, 0).text, meta.cell(1, 1).text = "Revision", "C"
    d.add_heading("1. States", level=1)
    d.add_paragraph("The pump has three states.")
    t = d.add_table(rows=3, cols=3)
    for i, row in enumerate([("State", "Pump", "Exit"), ("IDLE", "off", "START"),
                             ("RUN", "on", "high level")]):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    d.add_heading("2. Interlocks", level=1)
    d.add_paragraph("The pump never runs dry.", style="List Bullet")
    p = tmp_path / "note.docx"
    d.save(p)
    return p


def test_docx_paragraphs_carry_their_section(docx_file, ctx):
    r = read_source(docx_file, "src_doc", ctx)
    assert r.status == "read"
    c = next(c for c in r.chunks if c.text == "The pump never runs dry.")
    assert "2. Interlocks" in c.locator and "paragraph" in c.locator


def test_docx_table_rows_keep_headers_in_body_order(docx_file, ctx):
    r = read_source(docx_file, "src_doc", ctx)
    rows = [c for c in r.chunks if c.kind == "table_row"]
    run = next(c for c in rows if c.fields.get("State") == "RUN")
    assert run.text == "State: RUN | Pump: on | Exit: high level"
    assert "table 2, row 3" in run.locator and "1. States" in run.locator
    order = texts(r)
    assert order.index("The pump has three states.") < order.index(run.text) \
        < order.index("The pump never runs dry.")


def test_docx_two_column_table_reads_as_key_value(docx_file, ctx):
    r = read_source(docx_file, "src_doc", ctx)
    assert "Document: PCN-07" in texts(r)
    assert "Revision: C" in texts(r)


# --- xlsx ---------------------------------------------------------------------------------

@pytest.fixture
def xlsx_file(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parameters"
    ws.append(["Parameter register - collected from several revisions"])
    ws.append([])
    ws.append(["Parameter", "Tag", "Value", "Units", "Source", "Status"])
    ws.append(["High limit", "P-1", 0.9, "m", "CR-9", "Approved"])
    ws.append([])
    ws.append(["Low limit", "P-1", 0.1, "m", None, "Current"])
    log = wb.create_sheet("Hidden_History")
    log.append(["Record", "Change Summary", "Approved By"])
    log.append(["CR-9", "High limit 0.85 -> 0.9 m", "Board"])
    p = tmp_path / "register.xlsx"
    wb.save(p)
    return p


def test_xlsx_row_names_sheet_row_and_headers(xlsx_file, ctx):
    r = read_source(xlsx_file, "src_reg", ctx)
    row = next(c for c in r.chunks if c.fields.get("Source") == "CR-9")
    assert row.kind == "table_row"
    assert row.locator == "sheet Parameters, row 4"
    assert row.text == ("Parameter: High limit | Tag: P-1 | Value: 0.9 | Units: m | "
                        "Source: CR-9 | Status: Approved")


def test_xlsx_empty_cells_are_left_out_and_rows_keep_their_number(xlsx_file, ctx):
    r = read_source(xlsx_file, "src_reg", ctx)
    low = next(c for c in r.chunks if c.fields.get("Parameter") == "Low limit")
    assert low.locator == "sheet Parameters, row 6"
    assert "Source" not in low.fields


def test_xlsx_title_row_is_prose_and_every_sheet_is_read(xlsx_file, ctx):
    r = read_source(xlsx_file, "src_reg", ctx)
    title = next(c for c in r.chunks if c.locator == "sheet Parameters, row 1")
    assert title.kind == "prose"
    assert any(c.locator == "sheet Hidden_History, row 2" for c in r.chunks)


# --- eml ----------------------------------------------------------------------------------

THREAD = (
    "From: newest@x.example\n"
    "To: team@x.example\n"
    "Date: Thu, 12 Mar 2026 09:14:00 +0000\n"
    "Subject: RE: pump limit\n"
    "Content-Type: text/plain; charset=\"utf-8\"\n"
    "Content-Transfer-Encoding: quoted-printable\n"
    "MIME-Version: 1.0\n\n"
    "Approved. Use 0.9 m as the effective limit from now on, please update the mod=\n"
    "el.\n\n"
    "-----Original Message-----\n"
    "From: middle@x.example\n"
    "Sent: Tuesday, March 10, 2026 16:42\n"
    "Subject: RE: pump limit\n\n"
    "We could accept 0.9 m.\n\n"
    "-----Original Message-----\n"
    "From: oldest@x.example\n"
    "Sent: Monday, February 16, 2026 11:05\n"
    "Subject: pump limit\n\n"
    "My model has limit=3D0.85 m.\n"
)


@pytest.fixture
def eml_file(tmp_path):
    p = tmp_path / "thread.eml"
    p.write_bytes(THREAD.encode("utf-8"))
    return p


def test_eml_splits_messages_newest_first_with_sender_and_date(eml_file, ctx):
    r = read_source(eml_file, "src_mail", ctx)
    msgs = [c for c in r.chunks if c.kind == "email_message"]
    assert len(msgs) == 3
    assert [m.text.splitlines()[0] for m in msgs] == [
        "From: newest@x.example", "From: middle@x.example", "From: oldest@x.example"]
    assert "Date: Tuesday, March 10, 2026 16:42" in msgs[1].text
    assert "middle@x.example" in msgs[1].locator


def test_eml_is_quoted_printable_decoded(eml_file, ctx):
    r = read_source(eml_file, "src_mail", ctx)
    joined = "\n".join(texts(r))
    assert "update the model." in joined
    assert "limit=0.85 m" in joined and "=3D" not in joined


def test_eml_stale_value_is_in_the_message_that_said_it(eml_file, ctx):
    r = read_source(eml_file, "src_mail", ctx)
    old = next(c for c in r.chunks if "0.85" in c.text)
    assert "oldest@x.example" in old.text


def test_eml_orders_by_date_not_by_position(tmp_path, ctx):
    text = ("From: first@x.example\nDate: Mon, 2 Feb 2026 09:00:00 +0000\nSubject: s\n\n"
            "Early message.\n\n---- reply ----\nFrom: second@x.example\n"
            "Date: Tue, 3 Feb 2026 09:00:00 +0000\n\nLater message.\n")
    p = tmp_path / "t.eml"
    p.write_text(text, encoding="utf-8")
    r = read_source(p, "src_t", ctx)
    assert [c.text.splitlines()[0] for c in r.chunks] == ["From: second@x.example",
                                                         "From: first@x.example"]


@pytest.mark.parametrize("separator", ["-----Original Message-----", "---- Follow-up ----",
                                       "---- Original message ----", "-- forwarded reply --"])
def test_eml_separator_styles(tmp_path, ctx, separator):
    text = ("From: a@x.example\nDate: Mon, 2 Feb 2026 09:00:00 +0000\nSubject: s\n\nFirst.\n\n"
            f"{separator}\nFrom: b@x.example\nDate: Tue, 3 Feb 2026 09:00:00 +0000\n\nSecond.\n")
    p = tmp_path / "s.eml"
    p.write_text(text, encoding="utf-8")
    r = read_source(p, "src_s", ctx)
    assert len(r.chunks) == 2 and "Second." in r.chunks[0].text


def test_eml_angle_quoted_block_is_its_own_message(tmp_path, ctx):
    text = ("From: a@x.example\nDate: Tue, 3 Feb 2026 09:00:00 +0000\nSubject: s\n\n"
            "Agreed, 5 s.\n\nOn Mon, 2 Feb 2026 at 08:00, b@x.example wrote:\n"
            "> Use 4 s wait.\n> Thanks\n")
    p = tmp_path / "q.eml"
    p.write_text(text, encoding="utf-8")
    r = read_source(p, "src_q", ctx)
    assert len(r.chunks) == 2
    quoted = next(c for c in r.chunks if "Use 4 s wait." in c.text)
    assert "Agreed" not in quoted.text and ">" not in quoted.text
    assert "b@x.example" in quoted.text


# --- text and markdown --------------------------------------------------------------------

def test_text_paragraphs_have_line_ranges(tmp_path, ctx):
    p = tmp_path / "notes.txt"
    p.write_text("FIELD NOTES\n\n09:00 - pump started.\n09:05 - pump stopped.\n\nEnd.\n",
                 encoding="utf-8")
    r = read_source(p, "src_n", ctx)
    assert r.format == "text"
    c = next(c for c in r.chunks if "pump started" in c.text)
    assert c.locator == "lines 3-4"
    assert c.kind == "prose"


def test_markdown_items_carry_their_heading(tmp_path, ctx):
    p = tmp_path / "minutes.md"
    p.write_text("# Review Minutes\n\n**Date:** 2026-02-19\n\n## Decisions\n"
                 "- **D-01:** Pump first.\n- **D-02:** Valve second.\n\n"
                 "| Item | Owner |\n|---|---|\n| A-1 | Controls |\n", encoding="utf-8")
    r = read_source(p, "src_md", ctx)
    assert r.title == "Review Minutes"
    d2 = next(c for c in r.chunks if "D-02" in c.text)
    assert "D-01" not in d2.text
    assert d2.locator == "section 'Decisions', line 7"
    row = next(c for c in r.chunks if c.kind == "table_row")
    assert row.text == "Item: A-1 | Owner: Controls"


# --- code ---------------------------------------------------------------------------------

def test_modelica_is_code_with_comments_kept(tmp_path, ctx):
    p = tmp_path / "legacy.mo"
    p.write_text("model Old\n  // STALE: predates the change record.\n  parameter Real h = 0.85;\n"
                 "\nequation\n  der(x) = 1;\nend Old;\n", encoding="utf-8")
    r = read_source(p, "src_mo", ctx)
    assert r.format == "modelica"
    assert all(c.kind == "code" for c in r.chunks)
    first = r.chunks[0]
    assert "// STALE" in first.text and first.locator == "lines 1-3"


def test_plantuml_is_code(tmp_path, ctx):
    p = tmp_path / "arch.puml"
    p.write_text("@startuml\n' draft, incomplete\nA --> B : liquid\n@enduml\n", encoding="utf-8")
    r = read_source(p, "src_pu", ctx)
    assert r.format == "plantuml"
    assert "' draft, incomplete" in r.chunks[0].text and r.chunks[0].kind == "code"


# --- json ---------------------------------------------------------------------------------

def test_json_is_split_by_path(tmp_path, ctx):
    p = tmp_path / "export.json"
    p.write_text(json.dumps({"info": {"units": "SI"},
                             "items": [{"tag": "P-1", "x": 1.5}, {"tag": "P-2", "x": 2}]}),
                 encoding="utf-8")
    r = read_source(p, "src_js", ctx)
    locs = [c.locator for c in r.chunks]
    assert locs == ["$.info", "$.items"]
    assert '"P-1"' in r.chunks[1].text


def test_large_json_nodes_are_split_further(tmp_path, ctx):
    items = [{"tag": f"P-{i}", "note": "x" * 120} for i in range(12)]
    p = tmp_path / "big.json"
    p.write_text(json.dumps({"items": items}), encoding="utf-8")
    r = read_source(p, "src_big", ctx)
    assert [c.locator for c in r.chunks][:2] == ["$.items[0]", "$.items[1]"]
    assert r.chunks[0].fields == {"tag": "P-0", "note": "x" * 120}


# --- csv ----------------------------------------------------------------------------------

def test_csv_is_summarised_not_inlined(tmp_path, ctx):
    rows = ["time_s,level_m,state"] + [f"{t},{0.1 + t / 100:.2f},{'RUN' if t else 'IDLE'}"
                                       for t in range(0, 51)]
    p = tmp_path / "run.csv"
    p.write_text("\n".join(rows) + "\n", encoding="utf-8")
    r = read_source(p, "src_csv", ctx)
    assert all(c.kind == "data_summary" for c in r.chunks)
    summary = r.chunks[0]
    assert "rows: 51" in summary.text
    assert "time span: time_s from 0 to 50" in summary.text
    assert "columns: time_s, level_m, state" in summary.text
    level = next(c for c in r.chunks if "level_m" in c.locator)
    assert "min 0.1" in level.text and "max 0.6" in level.text
    assert "first 0.1" in level.text and "last 0.6" in level.text
    state = next(c for c in r.chunks if "state" in c.locator)
    assert "IDLE" in state.text and "RUN" in state.text
    assert len(r.chunks) == 4  # never one chunk per data row


# --- image --------------------------------------------------------------------------------

def test_image_with_vision_off_is_unread_with_reason(tmp_path, ctx):
    p = tmp_path / "diagram.png"
    p.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 16)
    r = read_source(p, "src_img", ctx)
    assert r.status == "unread"
    assert "vision" in r.reason
    assert r.chunks == []


class FakeVision:
    def __init__(self):
        self.calls = []

    def complete(self, *, prompt, input_text, schema, image=None, image_media_type=None):
        self.calls.append((prompt, image, image_media_type))
        return schema.model_validate({
            "labels": ["tank1", "valve1 (XV-1)"],
            "connections": [{"source": "valve1", "target": "tank1", "label": "liquid"}],
        })


def test_image_with_vision_on_becomes_diagram_text(tmp_path):
    settings = load_settings({"SPECALIVE_CACHE_DIR": str(tmp_path / "c"), "SPECALIVE_VISION": "1"})
    fake = FakeVision()
    p = tmp_path / "diagram.png"
    p.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 16)
    r = read_source(p, "src_img", ReadContext(settings=settings, llm=fake))
    assert fake.calls and fake.calls[0][1] == p.read_bytes()
    assert fake.calls[0][2] == "image/png"
    assert r.status == "partial" and "vision" in r.reason
    assert all(c.kind == "diagram_text" for c in r.chunks)
    joined = "\n".join(texts(r))
    assert "valve1 (XV-1)" in joined and "valve1 -> tank1 : liquid" in joined
